# -*- coding: utf-8 -*-
"""md→docx 后处理：对数模论文 docx 应用最终排版规范。

用法: python apply_docx_layout.py input.docx [output.docx]

前置约定（md 标题层级 -> Word 样式）:
  "#"   论文标题      -> Heading 1 -> 黑体三号(16pt) 居中
  "##"  一级标题(一、)  -> Heading 2 -> 黑体四号(14pt) 居中, 段前段后6磅 (含"摘要"标题)
  "###" 二级标题(1.1)  -> Heading 3 -> 黑体小四(12pt) 左对齐, 段后6磅
  "####" 三级标题(1.1.1)-> Heading 4 -> 黑体小四(12pt) 左对齐
  其余编号(1./（1）等) 用正文文本, 不加标题样式。

功能:
  1) 页面: A4(210x297mm), 上下72pt, 左右76.5pt
  2) 样式: 标题/摘要/关键词/一级~三级标题/正文/图表标题
  3) 内联字体分类: 中文宋体小四正体; 数学变量 TNR 小四斜体;
     名称/缩写/函数/算法/常量/单位 正体; 数学符号 Cambria Math 斜体
  4) 关键词行: "关键词："黑体小四, 内容宋体小四, 整行加粗
  5) 表格三线表、居中、内容居中宋体（默认五号、紧凑边距）
  6) 图片段落居中、无缩进
  7) 独立公式(oMathPara)居中, 并按顺序在右端加编号 (1),(2),...
"""
import argparse, re
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CN = "宋体"; LATIN = "Times New Roman"; MATH = "Cambria Math"; HEI = "黑体"
HEADING_STYLES = {"Title","Heading 1","Heading 2","Heading 3","Heading 4","Heading 5","Caption"}

# ---------- 内联字体分类 ----------
CJK_RE = re.compile(r'[\u3000-\u303f\u2013-\u201f\u2026\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uff00-\uffef]')
GREEK = "αβγδεζηθικλμνξοπρστυφχψωΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡΣΤΥΦΧΨΩ"
SPECIES_RE = re.compile(r'(?:Species|SSpecies)_[A-Za-z0-9_]+')
OP_CHARS = set("×=∈∑∏·∼≈→±≤≥≠⊂⊃∂∞")
LABELS = {
 "Species","SSpecies","Xylem","Bark","Diameter","Tissue","Component","Block","Subject",
 "Direction","Time","Treatment","VIF","RMSE","MAE","AIC","BIC","SE","MR","LR","HPL","LS",
 "OLS","DBSCAN","SPSS","VPA","SVG","II","L1","L2","Z","L",
 "Olson","Elastic","Net","LASSO","Ridge","Bootstrap","Spearman","GroupKFold","Leave","One","Out",
 "Marginal","Conditional","Akaike","log","ln","exp","sin","cos","tan","alpha","ratio","base","score",
 "code","problem","predictive","model","py","cv","results","csv","lrr","derived","data","processed","cleaned",
 # 常见医学/统计缩写与名称(正体)
 "NIPT","ROC","GC","DNA","AB","IVF","AUC","CV","CI","AE","GBDT","Logistic","Pearson","Turnbull",
 "T13","T18","T21","Spearman","OLS","GLM","ANOVA",
}
# 已知多字母数学变量(斜体); 其余多字母词默认按正体(名称/缩写/英文单词)处理
VAR_NAMES = {
 "LRR","RII","RCI","CN","TissueBark","LigninN","rs","us","sd","Gd","Dk",
 "kW","kXylem","kBark","ksingle","Srate","Emass","Ctotal","Wmass","WC","WN",
 "xylem","bark","total","single","rate","mass","Mxylem","Mbark","delta","lambda",
}
UNITS = {"g","kg","mg","mm","cm","m","km","h","min","mol","mmol","mL","L","℃"}

def classify_letter(ch):
    if ch == "g": return "unit"
    if ch in ("Z","L"): return "label"
    return "var"

def classify_word(word):
    if word in LABELS: return "label"
    if word in UNITS: return "unit"
    if word in VAR_NAMES: return "var"
    return "label"  # 未知多字母词默认正体(英文单词/缩写), 避免误斜体

def classify(text):
    out = []; i = 0; n = len(text)
    while i < n:
        ch = text[i]
        if CJK_RE.match(ch):
            j = i+1
            while j < n and CJK_RE.match(text[j]): j += 1
            out.append((text[i:j], "cjk")); i = j; continue
        m = SPECIES_RE.match(text, i)
        if m:
            out.append((m.group(0), "label")); i = m.end(); continue
        m = re.match(r"T(?:13|18|21)", text[i:])
        if m:
            out.append((m.group(0), "label")); i += m.end(); continue  # m.end() 为切片内偏移
        if ch in GREEK:
            out.append((ch, "greek" if ch != "π" else "greek_up")); i += 1; continue
        if ch.isdigit() or (ch in "+-" and i+1 < n and text[i+1].isdigit()):
            j = i+1
            while j < n and (text[j].isdigit() or text[j] in ".,%‰"): j += 1
            out.append((text[i:j], "num")); i = j; continue
        if ("a" <= ch <= "z") or ("A" <= ch <= "Z"):
            j = i+1
            while j < n and (("a" <= text[j] <= "z") or ("A" <= text[j] <= "Z")): j += 1
            word = text[i:j]
            out.append((word, classify_letter(word) if len(word) == 1 else classify_word(word)))
            i = j; continue
        if ch.isspace() or ch == "\u200b":
            out.append((ch, "ws")); i += 1; continue
        if ch in "-−~～":
            out.append((ch, "dash")); i += 1; continue
        if ch in "<>":
            out.append((ch, "rel")); i += 1; continue
        if ch in OP_CHARS:
            out.append((ch, "op")); i += 1; continue
        out.append((ch, "punct")); i += 1
    return out

def _set_rfonts(rPr, ascii_f, ea_f):
    """设置字体并清除主题属性(asciiTheme/eastAsiaTheme等)，避免主题字体覆盖显式字体。"""
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    for a in ("w:ascii","w:hAnsi","w:eastAsia","w:cs",
              "w:asciiTheme","w:hAnsiTheme","w:eastAsiaTheme","w:cstheme","w:hint"):
        rFonts.attrib.pop(qn(a), None)
    rFonts.set(qn("w:ascii"), ascii_f); rFonts.set(qn("w:hAnsi"), ascii_f)
    rFonts.set(qn("w:eastAsia"), ea_f)

def set_run_font(run, kind, bold=None, italic=None, sub=False, sup=False, size_pt=12):
    rPr = run._r.get_or_add_rPr()
    if kind == "cjk":
        a, e, it = CN, CN, False
    elif kind == "hei":
        a, e, it = HEI, HEI, False
    elif kind == "var":
        a, e, it = LATIN, CN, True
    elif kind in ("greek","greek_up","op"):
        a, e, it = MATH, CN, (kind in ("greek","op"))
    else:
        a, e, it = LATIN, CN, False
    _set_rfonts(rPr, a, e)
    for tag in ("w:i","w:iCs"):
        el = rPr.find(qn(tag))
        if el is not None: rPr.remove(el)
    if it:
        rPr.append(OxmlElement("w:i")); rPr.append(OxmlElement("w:iCs"))
    if bold is not None:
        for tag in ("w:b","w:bCs"):
            el = rPr.find(qn(tag))
            if el is not None: rPr.remove(el)
        if bold:
            rPr.append(OxmlElement("w:b")); rPr.append(OxmlElement("w:bCs"))
    for tag in ("w:sz","w:szCs"):
        el = rPr.find(qn(tag))
        if el is not None: rPr.remove(el)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size_pt*2))); rPr.append(sz)
    szCs = OxmlElement("w:szCs"); szCs.set(qn("w:val"), str(int(size_pt*2))); rPr.append(szCs)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "000000"); rPr.append(col)
    for tag in ("w:vertAlign",):
        el = rPr.find(qn(tag))
        if el is not None: rPr.remove(el)
    if sub:
        va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "subscript"); rPr.append(va)
    if sup:
        va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript"); rPr.append(va)

def has_nontext(run):
    r = run._r
    return (r.find(qn("w:drawing")) is not None or r.find(qn("w:object")) is not None
            or r.find(qn("w:pict")) is not None or r.find(qn("w:tab")) is not None
            or r.find(qn("w:br")) is not None or r.find(qn("m:oMath")) is not None)

def apply_inline(paragraph, size_pt=12):
    """把段落内纯文本 run 按内容分类重设字体；跳过图片/OMML/制表符/换行 run。"""
    from docx.text.run import Run
    for run in list(paragraph.runs):
        if has_nontext(run):
            continue
        text = run.text
        if text == "":
            continue
        bold = bool(run.font.bold)
        rPr = run._r.rPr
        va = None
        if rPr is not None:
            v = rPr.find(qn("w:vertAlign"))
            if v is not None: va = v.get(qn("w:val"))
        sub = (va == "subscript"); sup = (va == "superscript")
        r = run._r; parent = r.getparent()
        for chunk, kind in classify(text):
            new_r = OxmlElement("w:r"); new_r.append(OxmlElement("w:rPr"))
            wrapped = Run(new_r, run._parent)
            wrapped.text = chunk
            set_run_font(wrapped, kind, bold=bold, sub=sub, sup=sup, size_pt=size_pt)
            r.addprevious(new_r)
        parent.remove(r)

# ---------- 页面 ----------
def set_page(doc):
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Pt(72); sec.bottom_margin = Pt(72)
    sec.left_margin = Pt(76.5); sec.right_margin = Pt(76.5)

# ---------- 样式 ----------
def set_style_font(style, ascii_f, ea_f, size_pt, bold=None, italic=None):
    style.font.name = ascii_f
    style.font.size = Pt(size_pt)
    _set_rfonts(style.element.get_or_add_rPr(), ascii_f, ea_f)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic

def apply_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    set_style_font(normal, LATIN, CN, 12, bold=False, italic=False)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.25  # 正文 1.25 倍行距
    pf.first_line_indent = Pt(24)
    for bname in ("Body Text", "First Paragraph", "Compact", "Body Text 2"):
        if bname in st:
            st[bname].paragraph_format.line_spacing = 1.25
    # 论文标题: 黑体三号16pt 居中 (pandoc 中 "# 标题" -> Heading 1, 或 Title)
    for name in ("Heading 1", "Title"):
        if name in st:
            t = st[name]; set_style_font(t, HEI, HEI, 16, bold=False)
            t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t.paragraph_format.space_before = Pt(0); t.paragraph_format.space_after = Pt(0)
            t.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 一级标题(含摘要): 黑体四号14pt 居中 段前段后6磅
    if "Heading 2" in st:
        h = st["Heading 2"]; set_style_font(h, HEI, HEI, 14, bold=False)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.left_indent = Pt(0); h.paragraph_format.first_line_indent = Pt(0)
        h.paragraph_format.space_before = Pt(6); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 二级标题: 黑体小四12pt 左对齐 段后6磅
    if "Heading 3" in st:
        h = st["Heading 3"]; set_style_font(h, HEI, HEI, 12, bold=False)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.left_indent = Pt(0); h.paragraph_format.first_line_indent = Pt(0)
        h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 三级标题: 黑体小四12pt 左对齐顶格
    if "Heading 4" in st:
        h = st["Heading 4"]; set_style_font(h, HEI, HEI, 12, bold=False)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.left_indent = Pt(0); h.paragraph_format.first_line_indent = Pt(0)
        h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(0)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 图表标题: 宋体+TNR 五号10.5pt 加粗 居中 (自动创建 Caption)
    for cname in ("Caption", "Image Caption", "Table Caption"):
        if cname not in st:
            try:
                st.add_style(cname, 1)  # WD_STYLE_TYPE.PARAGRAPH
            except Exception:
                pass
        if cname in st:
            c = st[cname]; set_style_font(c, LATIN, CN, 10.5, bold=True)
            c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.first_line_indent = Pt(0)

def is_heading_para(p):
    try:
        return p.style.name in HEADING_STYLES
    except Exception:
        return False

def apply_abstract_and_keywords(doc):
    """摘要标题两空格; 关键词行特殊字体。"""
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if t in ("摘要", "摘 要", "摘  要") and is_heading_para(p):
            first = True
            for r in list(p.runs):
                if "摘" in r.text or "要" in r.text:
                    r.text = "摘\u3000\u3000要" if first else ""
                    first = False
            break
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("关键词") or t.startswith("关键字"):
            # 清空重建: "关键词："黑体小四; 其余宋体小四; 整行加粗
            from docx.text.run import Run
            m = re.match(r'^(关键词|关键字)\s*[:：]?\s*', t)
            if m:
                head = m.group(0)
                rest = t[m.end():]
                r = p.runs[0]._r if p.runs else None
                if r is not None:
                    parent = r.getparent()
                    for run in list(p.runs):
                        parent.remove(run._r)
                    nr1 = OxmlElement("w:r"); nr1.append(OxmlElement("w:rPr"))
                    w1 = Run(nr1, p); w1.text = head
                    set_run_font(w1, "hei", bold=True, size_pt=12)
                    parent.append(nr1)
                    nr2 = OxmlElement("w:r"); nr2.append(OxmlElement("w:rPr"))
                    w2 = Run(nr2, p); w2.text = rest
                    set_run_font(w2, "cjk", bold=True, size_pt=12)
                    parent.append(nr2)
            break

# ---------- 段落 ----------
CAPTION_NAMES = ("Caption", "Image Caption", "Table Caption")

def apply_paragraphs(doc):
    for p in doc.paragraphs:
        if is_heading_para(p):
            continue
        if p.text.strip().startswith("关键词") or p.text.strip().startswith("关键字"):
            continue  # 由 apply_abstract_and_keywords 处理
        if p.style.name in CAPTION_NAMES:
            apply_inline(p, size_pt=10.5)
            for r in p.runs:
                r.font.bold = True
            continue
        apply_inline(p, size_pt=12)

# ---------- 表格 ----------
def _set_tbl_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tblPr.append(borders)
    def edge(tag, sz):
        el = borders.find(qn("w:" + tag))
        if el is None:
            el = OxmlElement("w:" + tag); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), "auto")
    edge("top", 12); edge("bottom", 12)
    for t in ("left","right","insideH","insideV"):
        el = borders.find(qn("w:" + t))
        if el is None:
            el = OxmlElement("w:" + t); borders.append(el)
        el.set(qn("w:val"), "none"); el.set(qn("w:sz"), "0"); el.set(qn("w:space"), "0")
    if len(tbl.rows) > 0:
        for cell in tbl.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcB = tcPr.find(qn("w:tcBorders"))
            if tcB is None:
                tcB = OxmlElement("w:tcBorders"); tcPr.append(tcB)
            bottom = tcB.find(qn("w:bottom"))
            if bottom is None:
                bottom = OxmlElement("w:bottom"); tcB.append(bottom)
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "0"); bottom.set(qn("w:color"), "auto")


def _set_tbl_cell_margins(tbl, left=50, right=50, top=0, bottom=0):
    """设置表格单元格边距（twips），用于紧凑排版。"""
    tblPr = tbl._tbl.tblPr
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar"); tblPr.append(mar)
    for tag, val in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        el = mar.find(qn("w:" + tag))
        if el is None:
            el = OxmlElement("w:" + tag); mar.append(el)
        el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")


def _clear_first_line_chars(paragraph):
    """清除段落首行缩进（firstLine 与 firstLineChars 都要置 0）。

    坑点：Normal 样式常设 firstLineChars=200（2 字符首行缩进）。单元格段落若只清
    firstLine=0 而不清 firstLineChars=0，Word 仍按 2 字符缩进排版，导致表头/数字拆行。
    """
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind"); pPr.append(ind)
    ind.set(qn("w:firstLine"), "0"); ind.set(qn("w:firstLineChars"), "0")

def apply_tables(doc):
    for tbl in doc.tables:
        _set_tbl_borders(tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        # 紧凑单元格边距：左右 50 twips（约 0.09cm），上下 0
        _set_tbl_cell_margins(tbl, left=50, right=50, top=0, bottom=0)
        for row in tbl.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Pt(0)
                    _clear_first_line_chars(p)  # 必须同时清 firstLineChars，否则继承 Normal 的 2 字符首行缩进导致表头拆行
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE  # 三线表单倍行距
                    apply_inline(p, size_pt=10.5)  # 三线表默认五号

def style_captions(doc):
    """普通文本图表标题: 以 图N/表N 开头且较短(≤80字) -> 五号加粗居中。"""
    cap_re = re.compile(r'^(图|表)\s*\d')
    for p in doc.paragraphs:
        if is_heading_para(p):
            continue
        t = p.text.strip()
        if not t or len(t) > 80:
            continue
        if cap_re.match(t):
            apply_inline(p, size_pt=10.5)
            for r in p.runs:
                r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

# ---------- 图片居中 ----------
def center_images(doc):
    for p in doc.paragraphs:
        if is_heading_para(p):
            continue
        for r in p.runs:
            if r._r.find(qn("w:drawing")) is not None:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                break

# ---------- 公式居中+右编号 ----------
def number_formulas(doc):
    """独立公式(oMathPara)：公式居中(居中制表位)、右端编号 (N) 右对齐(右制表位)。

    采用双制表位法: 段落左对齐 + 居中制表位(文本宽度1/2) + 右制表位(文本宽度)，
    公式前放一个 tab(跳到居中位)、编号前放一个 tab(跳到右端)。
    md 中请勿手写公式编号；行内公式不编号。
    """
    count = 0
    sec = doc.sections[0]
    text_w = sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips
    center_pos = int(text_w / 2); right_pos = int(text_w)
    for p in doc.paragraphs:
        om = p._p.find(qn("m:oMathPara"))
        if om is None:
            continue
        # 清除本脚本先前追加的编号；重复排版时保留公式，只重建一份编号。
        for old_run in list(p.runs):
            if (re.fullmatch(r"\t\(\d+\)", old_run.text)
                    and old_run._r in om.itersiblings()):
                p._p.remove(old_run._r)
        count += 1
        pPr = p._p.get_or_add_pPr()
        jc = pPr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc"); pPr.append(jc)
        jc.set(qn("w:val"), "left")
        p.paragraph_format.first_line_indent = Pt(0)
        tabs = pPr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs"); pPr.insert(0, tabs)
        for el in list(tabs):
            tabs.remove(el)
        t1 = OxmlElement("w:tab"); t1.set(qn("w:val"), "center"); t1.set(qn("w:pos"), str(center_pos)); tabs.append(t1)
        t2 = OxmlElement("w:tab"); t2.set(qn("w:val"), "right"); t2.set(qn("w:pos"), str(right_pos)); tabs.append(t2)
        # 公式前保证有一个 tab run(跳到居中制表位)
        prev = om.getprevious()
        if prev is None or prev.find(qn("w:tab")) is None:
            trun = OxmlElement("w:r"); trun.append(OxmlElement("w:tab"))
            om.addprevious(trun)
        run = p.add_run("\t(%d)" % count)
        run.font.name = LATIN
        run.font.size = Pt(12)
    return count

def force_upright(doc):
    """三级标题、图注、表注等不斜体：标题与图表标题段落的所有 run 去斜体。"""
    cap_re = re.compile(r'^(图|表)\s*\d')
    for p in doc.paragraphs:
        t = p.text.strip() or ""
        is_cap = (p.style.name in ("Caption", "Image Caption", "Table Caption")
                  or (cap_re.match(t) and len(t) < 80))
        if p.style.name.startswith("Heading") or is_cap:
            for r in p.runs:
                r.font.italic = False

def format_references(doc):
    """参考文献：首行不缩进，每条文献独立成段（左对齐）。"""
    in_ref = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "参考文献":
            in_ref = True
            continue
        if in_ref:
            if p.style.name.startswith("Heading"):
                break
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def force_black(doc):
    """全部文字颜色设为黑色（覆盖样式级颜色，如 pandoc 默认标题蓝色）。"""
    def set_black(run):
        rPr = run._r.get_or_add_rPr()
        c = rPr.find(qn("w:color"))
        if c is None:
            c = OxmlElement("w:color"); rPr.append(c)
        c.set(qn("w:val"), "000000")
    for p in doc.paragraphs:
        for r in p.runs:
            set_black(r)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_black(r)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output", nargs="?")
    args = ap.parse_args()
    out = args.output or args.input
    doc = Document(args.input)
    set_page(doc)
    apply_styles(doc)
    apply_abstract_and_keywords(doc)
    apply_paragraphs(doc)
    apply_tables(doc)
    center_images(doc)
    style_captions(doc)
    force_upright(doc)
    format_references(doc)
    force_black(doc)
    nf = number_formulas(doc)
    doc.save(out)
    print("saved:", out, "| formulas numbered:", nf)

if __name__ == "__main__":
    main()
